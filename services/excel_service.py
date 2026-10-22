import os
from datetime import datetime
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from database import EXPORTS_DIR

class ExcelService:
    @staticmethod
    def export_movements(movements, filepath=None):
        """Exporta lista de movimentações para uma planilha Excel profissional."""
        if not filepath:
            os.makedirs(EXPORTS_DIR, exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filepath = os.path.abspath(os.path.join(EXPORTS_DIR, f"relatorio_movimentacoes_{timestamp}.xlsx"))

        data = []
        for m in movements:
            data.append({
                "Data / Hora": m.get("data_hora", ""),
                "Tipo": m.get("tipo", ""),
                "Código EPI": m.get("epi_codigo", ""),
                "Nome do EPI": m.get("epi_nome", ""),
                "C.A.": m.get("epi_ca", ""),
                "Quantidade": m.get("quantidade", 0),
                "Unidade": m.get("epi_unidade", "UN"),
                "Funcionário": m.get("func_nome") or "Almoxarifado / N/A",
                "Matrícula": m.get("func_matricula") or "",
                "Cargo": m.get("func_cargo") or "",
                "Departamento": m.get("func_depto") or "",
                "Quem Entregou / Responsável": m.get("responsavel_entrega", ""),
                "Motivo": m.get("motivo", ""),
                "Observação": m.get("observacao", "")
            })

        df = pd.DataFrame(data)

        # Salva e formata com OpenPyXL
        with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Movimentações", index=False)
            workbook = writer.book
            worksheet = writer.sheets["Movimentações"]

            # Estilos de cabeçalho
            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Azul corporativo
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            border_thin = Border(
                left=Side(style='thin', color='D1D5DB'),
                right=Side(style='thin', color='D1D5DB'),
                top=Side(style='thin', color='D1D5DB'),
                bottom=Side(style='thin', color='D1D5DB')
            )

            for col_idx, col in enumerate(worksheet.iter_cols(min_row=1, max_row=1), start=1):
                cell = col[0]
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            # Estilo das linhas de dados
            zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
            for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, max_row=worksheet.max_row), start=2):
                fill = zebra_fill if row_idx % 2 == 0 else PatternFill(fill_type=None)
                for cell in row:
                    cell.font = Font(name="Segoe UI", size=10)
                    cell.border = border_thin
                    if fill.fill_type:
                        cell.fill = fill
                    if cell.column_letter in ['F', 'G']: # Quantidade / Unidade
                        cell.alignment = Alignment(horizontal="center")

            # Ajuste dinâmico de largura das colunas
            for col in worksheet.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        return filepath

    @staticmethod
    def export_inventory(epis, filepath=None, only_in_stock=True):
        """Exporta inventário de estoque para Excel com grupos e filtros."""
        if only_in_stock:
            epis = [e for e in epis if float(e.get("quantidade", 0) or 0) > 0]

        if not filepath:
            os.makedirs(EXPORTS_DIR, exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filepath = os.path.abspath(os.path.join(EXPORTS_DIR, f"inventario_estoque_{timestamp}.xlsx"))

        from services.pdf_service import PDFService
        from services.epi_service import EPIService

        def _format_estado(est_raw):
            e = (str(est_raw) or "").upper().strip()
            if "NOVO" in e:
                return "Novo"
            elif "BOM" in e or "REUTIL" in e or "USADO" in e:
                return "Usado / Bom"
            elif "RUIM" in e or "MANUT" in e or "REPAR" in e:
                return "Manutenção"
            elif "DESC" in e or "SUCATA" in e or "INSERV" in e:
                return "Descarte"
            return est_raw or "Novo"

        data = []
        for e in epis:
            pu = float(e.get("preco_unitario", 0.0) or 0.0)
            grp_raw = (e.get("grupo") or "").strip()
            if grp_raw and grp_raw.upper() not in ["GERAL", "NONE", "NULL"]:
                grp = grp_raw.upper()
            else:
                grp = "DIVERSOS / OUTROS"

            loc_full = e.get("localizacao") or ""
            if e.get("localizacao_secundaria"):
                loc_full = f"{loc_full} | {e['localizacao_secundaria']}" if loc_full else e["localizacao_secundaria"]

            stock_items = EPIService.get_epi_stock_items(e.get("id"))
            if stock_items:
                active_items = [s for s in stock_items if int(s.get("quantidade", 0) or 0) > 0] if only_in_stock else stock_items
                if active_items:
                    for s_it in active_items:
                        tam_g = str(s_it.get("tamanho", "") or "").strip() or "Único"
                        q_g = int(s_it.get("quantidade", 0) or 0)
                        tot_g = q_g * pu
                        st_g = "Zerado" if q_g == 0 else ("Baixo" if q_g <= e.get("estoque_minimo", 0) else "Normal")
                        est_lbl = _format_estado(s_it.get("estado", "NOVO"))
                        data.append({
                            "Grupo / Família": grp,
                            "Código": e.get("codigo", ""),
                            "Nome do EPI": e.get("nome", ""),
                            "C.A.": e.get("ca", ""),
                            "Validade C.A.": e.get("validade_ca") or "N/D",
                            "Fabricante": e.get("fabricante") or "",
                            "Tamanho": tam_g,
                            "Estado do Material": est_lbl,
                            "Unidade": e.get("unidade", "UN"),
                            "Estoque Físico": q_g,
                            "Preço Unitário (R$)": pu,
                            "Valor Total (R$)": tot_g,
                            "Estoque Mínimo": e.get("estoque_minimo", 0),
                            "Status": st_g,
                            "Localização": loc_full,
                            "Descrição": e.get("descricao") or ""
                        })
                    continue

            qtd = int(e.get("quantidade", 0) or 0)
            tot = qtd * pu
            status = "Zerado" if qtd == 0 else ("Baixo" if qtd <= e.get("estoque_minimo", 0) else "Normal")

            data.append({
                "Grupo / Família": grp,
                "Código": e.get("codigo", ""),
                "Nome do EPI": e.get("nome", ""),
                "C.A.": e.get("ca", ""),
                "Validade C.A.": e.get("validade_ca") or "N/D",
                "Fabricante": e.get("fabricante") or "",
                "Tamanho": e.get("tamanho") or "Único",
                "Estado do Material": "Novo",
                "Unidade": e.get("unidade", "UN"),
                "Estoque Físico": qtd,
                "Preço Unitário (R$)": pu,
                "Valor Total (R$)": tot,
                "Estoque Mínimo": e.get("estoque_minimo", 0),
                "Status": status,
                "Localização": loc_full,
                "Descrição": e.get("descricao") or ""
            })

        df = pd.DataFrame(data)

        with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Estoque de EPI", index=False)
            workbook = writer.book
            worksheet = writer.sheets["Estoque de EPI"]

            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

            for col in worksheet.iter_cols(min_row=1, max_row=1):
                col[0].fill = header_fill
                col[0].font = header_font
                col[0].alignment = Alignment(horizontal="center", vertical="center")

            for col in worksheet.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                worksheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        return filepath
