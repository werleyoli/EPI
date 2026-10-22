import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from database import get_connection, EXPORTS_DIR
from services.audit_service import AuditService
from datetime import datetime

class ImportService:
    @staticmethod
    def generate_employee_template() -> str:
        """Gera uma planilha modelo Excel para importação de funcionários."""
        filepath = os.path.join(EXPORTS_DIR, "Modelo_Importacao_Funcionarios.xlsx")
        os.makedirs(EXPORTS_DIR, exist_ok=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Funcionarios"

        headers = ["Matricula*", "Nome Completo*", "CPF", "Cargo*", "Departamento*", "Data Admissao (AAAA-MM-DD)", "Status"]
        sample_rows = [
            ["MAT-101", "Lucas Pereira Silva", "123.456.789-00", "Mecânico de Manutenção", "Manutenção", "2024-01-15", "Ativo"],
            ["MAT-102", "Ana Beatriz Costa", "234.567.890-11", "Operadora de Produção", "Produção", "2024-02-01", "Ativo"],
            ["MAT-103", "Rodrigo Santos Alves", "345.678.901-22", "Eletricista Industrial", "Manutenção", "2024-03-10", "Ativo"],
        ]

        header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in sample_rows:
            ws.append(row)

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

        wb.save(filepath)
        return filepath

    @staticmethod
    def generate_epi_template() -> str:
        """Gera uma planilha modelo Excel para importação de EPIs."""
        filepath = os.path.join(EXPORTS_DIR, "Modelo_Importacao_EPIs.xlsx")
        os.makedirs(EXPORTS_DIR, exist_ok=True)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "EPIs"

        headers = [
            "Codigo*", "Nome do EPI*", "CA*", "Validade CA (AAAA-MM-DD)", "Fabricante",
            "Tamanho", "Unidade", "Quantidade Inicial*", "Estoque Minimo*", "Preco Unitario (R$)",
            "Vida Util Dias", "Localizacao Almoxarifado"
        ]
        sample_rows = [
            ["EPI-101", "Capacete de Segurança Classe B", "41001", "2028-11-20", "3M", "Único", "UN", 50, 10, 38.50, 365, "Prateleira A1"],
            ["EPI-102", "Óculos de Proteção Incolor", "34110", "2027-08-15", "Kalipso", "Único", "UN", 100, 20, 14.90, 90, "Prateleira B2"],
            ["EPI-103", "Luva de Vaqueta Mista", "28540", "2028-05-10", "Danny", "G", "PAR", 60, 15, 26.00, 45, "Prateleira C3"],
        ]

        header_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for row in sample_rows:
            ws.append(row)

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

        wb.save(filepath)
        return filepath

    @staticmethod
    def import_employees_excel(file_path: str) -> dict:
        """Importa funcionários de uma planilha Excel (.xlsx) ou CSV."""
        if file_path.endswith(".csv"):
            df = pd.read_csv(file_path)
        else:
            df = pd.read_excel(file_path)

        # Normaliza nomes de colunas
        col_map = {}
        for col in df.columns:
            c_clean = str(col).lower().replace("*", "").replace(" ", "_")
            if "matr" in c_clean:
                col_map[col] = "matricula"
            elif "nome" in c_clean:
                col_map[col] = "nome"
            elif "cpf" in c_clean:
                col_map[col] = "cpf"
            elif "cargo" in c_clean or "funcao" in c_clean:
                col_map[col] = "cargo"
            elif "dep" in c_clean or "setor" in c_clean:
                col_map[col] = "departamento"
            elif "admis" in c_clean or "data" in c_clean:
                col_map[col] = "data_admissao"
            elif "status" in c_clean or "situacao" in c_clean:
                col_map[col] = "status"

        df = df.rename(columns=col_map)

        inseridos = 0
        atualizados = 0
        erros = []
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        with get_connection() as conn:
            cursor = conn.cursor()
            for idx, row in df.iterrows():
                linha_num = idx + 2
                mat = str(row.get("matricula", "")).strip()
                nome = str(row.get("nome", "")).strip()
                cargo = str(row.get("cargo", "")).strip()
                depto = str(row.get("departamento", "")).strip() or "Geral"

                if not mat or mat == "nan" or not nome or nome == "nan":
                    continue

                cpf = str(row.get("cpf", "")).strip()
                if cpf == "nan":
                    cpf = ""

                dt_adm = str(row.get("data_admissao", "")).strip()
                if dt_adm == "nan" or not dt_adm:
                    dt_adm = datetime.now().strftime("%Y-%m-%d")
                elif len(dt_adm) > 10:
                    dt_adm = dt_adm[:10]

                st = str(row.get("status", "Ativo")).strip()
                if st == "nan" or not st:
                    st = "Ativo"

                try:
                    cursor.execute("SELECT id FROM funcionarios WHERE matricula = ?;", (mat,))
                    existente = cursor.fetchone()
                    if existente:
                        cursor.execute("""
                            UPDATE funcionarios
                            SET nome = ?, cpf = ?, cargo = ?, departamento = ?, data_admissao = ?, status = ?
                            WHERE id = ?;
                        """, (nome, cpf, cargo, depto, dt_adm, st, existente["id"]))
                        atualizados += 1
                    else:
                        cursor.execute("""
                            INSERT INTO funcionarios (matricula, nome, cpf, cargo, departamento, data_admissao, status, created_at)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?);
                        """, (mat, nome, cpf, cargo, depto, dt_adm, st, now))
                        inseridos += 1
                except Exception as ex:
                    erros.append(f"Linha {linha_num} ({mat}): {str(ex)}")

            conn.commit()

        AuditService.log(
            acao="IMPORTACAO_EM_MASSA",
            modulo="FUNCIONARIOS",
            detalhes=f"Importação via Excel: {inseridos} novos colaboradores, {atualizados} atualizados. Arquivo: {os.path.basename(file_path)}"
        )

        return {"inseridos": inseridos, "atualizados": atualizados, "erros": erros}

    @staticmethod
    def import_epis_excel(file_path: str) -> dict:
        """Importa EPIs de uma planilha Excel (.xlsx) ou CSV."""
        if file_path.endswith(".csv"):
            df = pd.read_csv(file_path)
        else:
            df = pd.read_excel(file_path)

        col_map = {}
        for col in df.columns:
            c_clean = str(col).lower().replace("*", "").replace(" ", "_")
            if "cod" in c_clean:
                col_map[col] = "codigo"
            elif "nome" in c_clean or "equip" in c_clean or "epi" in c_clean:
                col_map[col] = "nome"
            elif "ca" in c_clean and "val" not in c_clean:
                col_map[col] = "ca"
            elif "val" in c_clean:
                col_map[col] = "validade_ca"
            elif "fab" in c_clean or "marca" in c_clean:
                col_map[col] = "fabricante"
            elif "tam" in c_clean:
                col_map[col] = "tamanho"
            elif "unid" in c_clean:
                col_map[col] = "unidade"
            elif "quant" in c_clean or "saldo" in c_clean:
                col_map[col] = "quantidade"
            elif "min" in c_clean:
                col_map[col] = "estoque_minimo"
            elif "prec" in c_clean or "custo" in c_clean or "valor" in c_clean:
                col_map[col] = "preco_unitario"
            elif "vida" in c_clean or "dias" in c_clean:
                col_map[col] = "dias_vida_util"
            elif "loc" in c_clean or "prat" in c_clean:
                col_map[col] = "localizacao"
            elif "grup" in c_clean or "famil" in c_clean:
                col_map[col] = "grupo"

        df = df.rename(columns=col_map)

        inseridos = 0
        atualizados = 0
        erros = []
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        with get_connection() as conn:
            cursor = conn.cursor()
            for idx, row in df.iterrows():
                linha_num = idx + 2
                cod = str(row.get("codigo", "")).strip()
                nome = str(row.get("nome", "")).strip()
                ca = str(row.get("ca", "")).strip()

                if not cod or cod == "nan" or not nome or nome == "nan":
                    continue

                def _clean_str(v, def_val=""):
                    s = str(v).strip()
                    return def_val if s == "nan" or not s else s

                def _clean_num(v, def_val=0, is_float=False):
                    try:
                        s = str(v).replace(",", ".").strip()
                        return float(s) if is_float else int(float(s))
                    except Exception:
                        return def_val

                val_ca = _clean_str(row.get("validade_ca"), "")
                if len(val_ca) > 10:
                    val_ca = val_ca[:10]

                fab = _clean_str(row.get("fabricante"), "Genérico")
                tam = _clean_str(row.get("tamanho"), "Único")
                unidade = _clean_str(row.get("unidade"), "UN").upper()
                qtd = _clean_num(row.get("quantidade"), 0)
                min_q = _clean_num(row.get("estoque_minimo"), 5)
                preco = _clean_num(row.get("preco_unitario"), 0.0, is_float=True)
                vida = _clean_num(row.get("dias_vida_util"), 90)
                loc = _clean_str(row.get("localizacao"), "Almoxarifado")
                grp = _clean_str(row.get("grupo"), "").strip().upper()

                try:
                    cursor.execute("SELECT id, grupo FROM epis WHERE codigo = ?;", (cod,))
                    existente = cursor.fetchone()
                    if existente:
                        # Se informado novo grupo na planilha, usa; senão mantém o que já estava
                        final_grp = grp if grp else (existente["grupo"] or "DIVERSOS / OUTROS")
                        cursor.execute("INSERT OR IGNORE INTO grupos_epi (nome, descricao, created_at) VALUES (?, ?, ?);",
                                       (final_grp, f"Grupo {final_grp}", now))
                        cursor.execute("""
                            UPDATE epis
                            SET nome = ?, ca = ?, validade_ca = ?, fabricante = ?, tamanho = ?,
                                unidade = ?, estoque_minimo = ?, preco_unitario = ?,
                                dias_vida_util = ?, localizacao = ?, grupo = ?
                            WHERE id = ?;
                        """, (nome, ca, val_ca, fab, tam, unidade, min_q, preco, vida, loc, final_grp, existente["id"]))
                        atualizados += 1
                    else:
                        final_grp = grp if grp else "DIVERSOS / OUTROS"
                        cursor.execute("INSERT OR IGNORE INTO grupos_epi (nome, descricao, created_at) VALUES (?, ?, ?);",
                                       (final_grp, f"Grupo {final_grp}", now))
                        cursor.execute("""
                            INSERT INTO epis (codigo, nome, ca, validade_ca, fabricante, tamanho,
                                            unidade, quantidade, estoque_minimo, preco_unitario,
                                            dias_vida_util, localizacao, grupo, created_at)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
                        """, (cod, nome, ca, val_ca, fab, tam, unidade, qtd, min_q, preco, vida, loc, final_grp, now))
                        inseridos += 1
                except Exception as ex:
                    erros.append(f"Linha {linha_num} ({cod}): {str(ex)}")

            conn.commit()

        AuditService.log(
            acao="IMPORTACAO_EM_MASSA",
            modulo="ESTOQUE",
            detalhes=f"Importação via Excel: {inseridos} novos EPIs cadastrados, {atualizados} atualizados. Arquivo: {os.path.basename(file_path)}"
        )

        return {"inseridos": inseridos, "atualizados": atualizados, "erros": erros}
